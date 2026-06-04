#!/usr/bin/env python3
"""Convert markdown thesis chapter to .docx (basic formatting)."""
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_paragraph(doc, text, style=None, bold=False):
    if not text.strip():
        return
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    if bold:
        run.bold = True
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    return p


def parse_md(path: Path) -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(line[3:].strip())
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            p.paragraph_format.space_before = Pt(12)
            i += 1
            continue
        if line.startswith("### "):
            add_paragraph(doc, line[4:].strip(), bold=True)
            i += 1
            continue
        if line.startswith("#### "):
            add_paragraph(doc, line[5:].strip(), bold=True)
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("**ВСТАВИТЬ РИСУНОК"):
            add_paragraph(doc, line.strip("*"), bold=True)
            i += 1
            if i < len(lines) and lines[i].startswith("*"):
                add_paragraph(doc, lines[i].strip("* ").strip())
                i += 1
            continue
        if line.startswith("|") and "|" in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-:|]+\|$", lines[i]):
                    table_lines.append(lines[i])
                i += 1
            if table_lines:
                cols = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                table = doc.add_table(rows=len(table_lines), cols=len(cols))
                table.style = "Table Grid"
                for ri, tl in enumerate(table_lines):
                    cells = [c.strip() for c in tl.split("|")[1:-1]]
                    for ci, cell in enumerate(cells):
                        table.rows[ri].cells[ci].text = cell
            continue
        if line.startswith("- ") or line.startswith("•\t"):
            bullet = line[2:].strip() if line.startswith("- ") else line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(bullet)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14)
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            if code_lines:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(11)
            continue
        if line.strip():
            add_paragraph(doc, line.strip())
        i += 1
    return doc


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent.parent / "ГЛАВА_3_РАЗРАБОТКА.md"
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    doc = parse_md(src)
    doc.save(out)
    print(f"Saved: {out}")
